#!/usr/bin/env python3
"""
Business Model Innovation Excel Sheet Generator

This script creates a comprehensive business model innovation worksheet
that can be used for strategic planning and innovation analysis.
"""

import os
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


def create_business_model_innovation_sheet():
    """Create a comprehensive business model innovation Excel worksheet."""
    
    # Create a new workbook
    wb = Workbook()
    
    # Remove the default sheet and create our custom sheets
    wb.remove(wb.active)
    
    # Create main sheets
    overview_sheet = wb.create_sheet("BMI Overview")
    canvas_sheet = wb.create_sheet("Business Model Canvas")
    innovation_sheet = wb.create_sheet("Innovation Framework")
    analysis_sheet = wb.create_sheet("Analysis & Metrics")
    
    # Style definitions
    header_font = Font(bold=True, size=14, color="FFFFFF")
    subheader_font = Font(bold=True, size=12)
    regular_font = Font(size=11)
    header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
    light_fill = PatternFill(start_color="E7F3FF", end_color="E7F3FF", fill_type="solid")
    border = Border(
        left=Side(style='thin'),
        right=Side(style='thin'),
        top=Side(style='thin'),
        bottom=Side(style='thin')
    )
    
    # Configure Overview Sheet
    setup_overview_sheet(overview_sheet, header_font, header_fill, regular_font, border)
    
    # Configure Business Model Canvas Sheet
    setup_canvas_sheet(canvas_sheet, header_font, subheader_font, header_fill, light_fill, regular_font, border)
    
    # Configure Innovation Framework Sheet
    setup_innovation_sheet(innovation_sheet, header_font, subheader_font, header_fill, light_fill, regular_font, border)
    
    # Configure Analysis & Metrics Sheet
    setup_analysis_sheet(analysis_sheet, header_font, subheader_font, header_fill, light_fill, regular_font, border)
    
    return wb


def setup_overview_sheet(sheet, header_font, header_fill, regular_font, border):
    """Set up the BMI Overview sheet."""
    
    # Title
    sheet['A1'] = "Business Model Innovation Overview"
    sheet['A1'].font = Font(bold=True, size=16)
    sheet.merge_cells('A1:G1')
    
    # Introduction section
    sheet['A3'] = "Purpose"
    sheet['A3'].font = header_font
    sheet['A3'].fill = header_fill
    sheet['A3'].border = border
    sheet.merge_cells('A3:G3')
    
    purpose_text = ("This workbook provides a comprehensive framework for business model innovation. "
                   "It includes tools for mapping current business models, identifying innovation "
                   "opportunities, and tracking progress toward new value propositions.")
    
    sheet['A4'] = purpose_text
    sheet['A4'].font = regular_font
    sheet['A4'].alignment = Alignment(wrap_text=True, vertical='top')
    sheet.merge_cells('A4:G6')
    sheet.row_dimensions[4].height = 45
    
    # Instructions section
    sheet['A8'] = "How to Use This Workbook"
    sheet['A8'].font = header_font
    sheet['A8'].fill = header_fill
    sheet['A8'].border = border
    sheet.merge_cells('A8:G8')
    
    instructions = [
        "1. Start with the 'Business Model Canvas' sheet to map your current business model",
        "2. Use the 'Innovation Framework' sheet to identify innovation opportunities",
        "3. Track progress and metrics in the 'Analysis & Metrics' sheet",
        "4. Regularly review and update all sections as your business model evolves"
    ]
    
    for i, instruction in enumerate(instructions, start=9):
        sheet[f'A{i}'] = instruction
        sheet[f'A{i}'].font = regular_font
        sheet.merge_cells(f'A{i}:G{i}')
    
    # Set column widths
    for col in range(1, 8):
        sheet.column_dimensions[get_column_letter(col)].width = 15


def setup_canvas_sheet(sheet, header_font, subheader_font, header_fill, light_fill, regular_font, border):
    """Set up the Business Model Canvas sheet."""
    
    # Title
    sheet['A1'] = "Business Model Canvas"
    sheet['A1'].font = Font(bold=True, size=16)
    sheet.merge_cells('A1:I1')
    
    # Canvas structure - 9 building blocks
    canvas_blocks = [
        ("Key Partners", "A3", "A3:B8"),
        ("Key Activities", "A9", "A9:B14"),
        ("Key Resources", "A15", "A15:B20"),
        ("Value Propositions", "C3", "C3:E14"),
        ("Customer Relationships", "F3", "F3:G8"),
        ("Channels", "F9", "F9:G14"),
        ("Customer Segments", "H3", "H3:I14"),
        ("Cost Structure", "A21", "A21:E26"),
        ("Revenue Streams", "F21", "F21:I26")
    ]
    
    for block_name, header_cell, merge_range in canvas_blocks:
        # Header
        sheet[header_cell] = block_name
        sheet[header_cell].font = header_font
        sheet[header_cell].fill = header_fill
        sheet[header_cell].border = border
        sheet[header_cell].alignment = Alignment(horizontal='center', vertical='center')
        
        # Merge header cells
        header_row = int(header_cell[1:])
        if block_name in ["Key Partners", "Key Activities", "Key Resources"]:
            sheet.merge_cells(f'A{header_row}:B{header_row}')
        elif block_name == "Value Propositions":
            sheet.merge_cells(f'C{header_row}:E{header_row}')
        elif block_name in ["Customer Relationships", "Channels"]:
            sheet.merge_cells(f'F{header_row}:G{header_row}')
        elif block_name == "Customer Segments":
            sheet.merge_cells(f'H{header_row}:I{header_row}')
        elif block_name == "Cost Structure":
            sheet.merge_cells(f'A{header_row}:E{header_row}')
        elif block_name == "Revenue Streams":
            sheet.merge_cells(f'F{header_row}:I{header_row}')
    
    # Add content cells with light fill
    content_areas = [
        ("A4:B8", "List your key partners and suppliers"),
        ("A10:B14", "Describe key activities your business performs"),
        ("A16:B20", "List key resources needed for your business"),
        ("C4:E14", "Describe your value propositions for customers"),
        ("F4:G8", "How do you build relationships with customers?"),
        ("F10:G14", "How do you reach and deliver to customers?"),
        ("H4:I14", "Who are your target customer segments?"),
        ("A22:E26", "What are your key costs?"),
        ("F22:I26", "How do you generate revenue?")
    ]
    
    for area, placeholder in content_areas:
        start_cell = area.split(':')[0]
        sheet[start_cell] = placeholder
        sheet[start_cell].font = Font(italic=True, color="666666")
        sheet[start_cell].alignment = Alignment(wrap_text=True, vertical='top')
        
        # Apply light fill to the entire area
        for row in sheet[area]:
            for cell in row:
                cell.fill = light_fill
                cell.border = border
    
    # Set column widths
    column_widths = [15, 15, 20, 20, 20, 15, 15, 15, 15]
    for i, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    
    # Set row heights
    for row_num in range(4, 27):
        sheet.row_dimensions[row_num].height = 25


def setup_innovation_sheet(sheet, header_font, subheader_font, header_fill, light_fill, regular_font, border):
    """Set up the Innovation Framework sheet."""
    
    # Title
    sheet['A1'] = "Business Model Innovation Framework"
    sheet['A1'].font = Font(bold=True, size=16)
    sheet.merge_cells('A1:F1')
    
    # Innovation Categories
    sheet['A3'] = "Innovation Type"
    sheet['B3'] = "Description"
    sheet['C3'] = "Current State"
    sheet['D3'] = "Opportunity"
    sheet['E3'] = "Priority"
    sheet['F3'] = "Timeline"
    
    for col in ['A3', 'B3', 'C3', 'D3', 'E3', 'F3']:
        sheet[col].font = header_font
        sheet[col].fill = header_fill
        sheet[col].border = border
        sheet[col].alignment = Alignment(horizontal='center')
    
    # Innovation types
    innovation_types = [
        ("Value Proposition", "New ways to create and deliver value"),
        ("Revenue Model", "New ways to monetize value"),
        ("Customer Segments", "New target markets or segments"),
        ("Channels", "New distribution or communication channels"),
        ("Key Resources", "New assets or capabilities"),
        ("Key Activities", "New processes or operations"),
        ("Key Partners", "New strategic partnerships"),
        ("Cost Structure", "New cost optimization approaches"),
        ("Technology", "New technological capabilities"),
        ("Platform", "Platform-based business models")
    ]
    
    for i, (innovation_type, description) in enumerate(innovation_types, start=4):
        sheet[f'A{i}'] = innovation_type
        sheet[f'B{i}'] = description
        sheet[f'C{i}'] = "Describe current state"
        sheet[f'D{i}'] = "Identify opportunities"
        sheet[f'E{i}'] = "High/Medium/Low"
        sheet[f'F{i}'] = "Q1/Q2/Q3/Q4"
        
        for col in ['A', 'B', 'C', 'D', 'E', 'F']:
            cell = sheet[f'{col}{i}']
            cell.font = regular_font if col == 'A' else Font(size=10, italic=True, color="666666")
            cell.fill = light_fill
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    
    # Set column widths
    column_widths = [20, 30, 25, 25, 15, 15]
    for i, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    
    # Set row heights
    for row_num in range(4, 14):
        sheet.row_dimensions[row_num].height = 30


def setup_analysis_sheet(sheet, header_font, subheader_font, header_fill, light_fill, regular_font, border):
    """Set up the Analysis & Metrics sheet."""
    
    # Title
    sheet['A1'] = "Business Model Analysis & Metrics"
    sheet['A1'].font = Font(bold=True, size=16)
    sheet.merge_cells('A1:E1')
    
    # Metrics section
    sheet['A3'] = "Key Performance Indicators"
    sheet['A3'].font = subheader_font
    sheet.merge_cells('A3:E3')
    
    sheet['A5'] = "Metric"
    sheet['B5'] = "Current Value"
    sheet['C5'] = "Target Value"
    sheet['D5'] = "Timeline"
    sheet['E5'] = "Status"
    
    for col in ['A5', 'B5', 'C5', 'D5', 'E5']:
        sheet[col].font = header_font
        sheet[col].fill = header_fill
        sheet[col].border = border
        sheet[col].alignment = Alignment(horizontal='center')
    
    # Sample metrics
    metrics = [
        "Customer Acquisition Cost",
        "Customer Lifetime Value",
        "Monthly Recurring Revenue",
        "Churn Rate",
        "Market Share",
        "Innovation Pipeline Value",
        "Time to Market",
        "Revenue per Customer",
        "Profit Margin",
        "Customer Satisfaction Score"
    ]
    
    for i, metric in enumerate(metrics, start=6):
        sheet[f'A{i}'] = metric
        sheet[f'B{i}'] = "Enter current value"
        sheet[f'C{i}'] = "Enter target value"
        sheet[f'D{i}'] = "Enter timeline"
        sheet[f'E{i}'] = "On Track/At Risk/Behind"
        
        for col in ['A', 'B', 'C', 'D', 'E']:
            cell = sheet[f'{col}{i}']
            cell.font = regular_font if col == 'A' else Font(size=10, italic=True, color="666666")
            cell.fill = light_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center' if col in ['D', 'E'] else 'left')
    
    # Risk Assessment section
    sheet['A17'] = "Risk Assessment"
    sheet['A17'].font = subheader_font
    sheet.merge_cells('A17:E17')
    
    sheet['A19'] = "Risk Factor"
    sheet['B19'] = "Impact"
    sheet['C19'] = "Probability"
    sheet['D19'] = "Mitigation Strategy"
    sheet['E19'] = "Owner"
    
    for col in ['A19', 'B19', 'C19', 'D19', 'E19']:
        sheet[col].font = header_font
        sheet[col].fill = header_fill
        sheet[col].border = border
        sheet[col].alignment = Alignment(horizontal='center')
    
    # Sample risks
    risks = [
        "Market disruption",
        "Technology obsolescence",
        "Competitive pressure",
        "Regulatory changes",
        "Economic downturn"
    ]
    
    for i, risk in enumerate(risks, start=20):
        sheet[f'A{i}'] = risk
        sheet[f'B{i}'] = "High/Medium/Low"
        sheet[f'C{i}'] = "High/Medium/Low"
        sheet[f'D{i}'] = "Describe mitigation approach"
        sheet[f'E{i}'] = "Assign owner"
        
        for col in ['A', 'B', 'C', 'D', 'E']:
            cell = sheet[f'{col}{i}']
            cell.font = regular_font if col == 'A' else Font(size=10, italic=True, color="666666")
            cell.fill = light_fill
            cell.border = border
            cell.alignment = Alignment(wrap_text=True, vertical='top')
    
    # Set column widths
    column_widths = [25, 20, 20, 30, 20]
    for i, width in enumerate(column_widths, start=1):
        sheet.column_dimensions[get_column_letter(i)].width = width
    
    # Set row heights for better visibility
    for row_num in range(6, 25):
        sheet.row_dimensions[row_num].height = 25


def main():
    """Main function to generate the business model innovation Excel file."""
    try:
        print("Generating Business Model Innovation Excel sheet...")
        
        # Create the workbook
        workbook = create_business_model_innovation_sheet()
        
        # Save the workbook
        output_path = os.path.join(os.path.dirname(__file__), "business-model-innovation.xlsx")
        workbook.save(output_path)
        
        print(f"Successfully created: {output_path}")
        print("\nThe Excel file contains the following sheets:")
        print("1. BMI Overview - Introduction and usage instructions")
        print("2. Business Model Canvas - Complete business model mapping")
        print("3. Innovation Framework - Innovation opportunities analysis")
        print("4. Analysis & Metrics - KPIs and risk assessment")
        
        return True
        
    except Exception as e:
        print(f"Error creating Excel file: {str(e)}")
        return False


if __name__ == "__main__":
    main()